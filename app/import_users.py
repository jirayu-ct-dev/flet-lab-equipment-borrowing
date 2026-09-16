from __future__ import annotations

import csv
import io
from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook, load_workbook

from app.errors import ValidationError
from app.service import AppService, USERNAME_RE


COLUMNS = (
    "username",
    "full_name",
    "backup_email",
    "role",
    "user_type",
    "faculty",
    "department",
    "cohort",
    "class_group",
)


@dataclass(frozen=True, slots=True)
class ImportRow:
    number: int
    values: dict[str, str]
    resolved: dict[str, object]
    errors: tuple[str, ...]


@dataclass(frozen=True, slots=True)
class ImportPreview:
    rows: tuple[ImportRow, ...]

    @property
    def valid_count(self) -> int:
        return sum(not row.errors for row in self.rows)

    @property
    def error_count(self) -> int:
        return len(self.rows) - self.valid_count


def template_bytes(extension: str) -> bytes:
    example = (
        "650000001",
        "สมชาย ใจดี",
        "somchai@example.com",
        "borrower",
        "student",
        "วิทยาศาสตร์",
        "วิทยาการคอมพิวเตอร์",
        "65",
        "1",
    )
    if extension.lower() == "csv":
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow(COLUMNS)
        writer.writerow(example)
        return output.getvalue().encode("utf-8-sig")
    output = io.BytesIO()
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "users"
    sheet.append(COLUMNS)
    sheet.append(example)
    workbook.save(output)
    return output.getvalue()


def _read_rows(name: str, data: bytes) -> list[dict[str, str]]:
    suffix = Path(name).suffix.lower()
    if suffix == ".csv":
        text = data.decode("utf-8-sig")
        return [{str(k).strip(): str(v or "").strip() for k, v in row.items()} for row in csv.DictReader(io.StringIO(text))]
    if suffix == ".xlsx":
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        try:
            headers = [str(value or "").strip() for value in next(rows)]
        except StopIteration:
            return []
        return [
            {headers[index]: str(value or "").strip() for index, value in enumerate(values) if index < len(headers)}
            for values in rows
            if any(value not in (None, "") for value in values)
        ]
    raise ValidationError("รองรับเฉพาะไฟล์ CSV และ XLSX")


def preview_users(service: AppService, name: str, data: bytes) -> ImportPreview:
    raw_rows = _read_rows(name, data)
    if not raw_rows:
        raise ValidationError("ไฟล์ไม่มีข้อมูลผู้ใช้")
    existing = {row["username"].casefold() for row in service.list_users()}
    seen: set[str] = set()
    faculties = {row["name"].strip().casefold(): row for row in service.list_master("faculty", active_only=True)}
    departments = {(row["faculty_id"], row["name"].strip().casefold()): row for row in service.list_master("department", active_only=True)}
    cohorts = {(row["department_id"], row["name"].strip().casefold()): row for row in service.list_master("cohort", active_only=True)}
    class_groups = {(row["cohort_id"], row["name"].strip().casefold()): row for row in service.list_master("class_group", active_only=True)}
    preview: list[ImportRow] = []
    for number, raw in enumerate(raw_rows, start=2):
        values = {column: raw.get(column, "").strip() for column in COLUMNS}
        values["role"] = values["role"] or "borrower"
        values["user_type"] = values["user_type"] or "student"
        errors: list[str] = []
        username = values["username"]
        key = username.casefold()
        if not USERNAME_RE.fullmatch(username):
            errors.append("username ไม่ถูกต้อง")
        if key in existing or key in seen:
            errors.append("username ซ้ำ")
        if not values["full_name"]:
            errors.append("ไม่มีชื่อ–นามสกุล")
        if values["role"] not in {"admin", "borrower"}:
            errors.append("role ต้องเป็น admin หรือ borrower")
        if values["user_type"] not in {"student", "teacher", "staff"}:
            errors.append("user_type ต้องเป็น student, teacher หรือ staff")

        resolved: dict[str, object] = {
            "username": username,
            "full_name": values["full_name"],
            "backup_email": values["backup_email"] or None,
            "role": values["role"],
            "user_type": values["user_type"],
        }
        if values["role"] == "borrower":
            faculty = faculties.get(values["faculty"].casefold())
            department = departments.get((faculty["id"], values["department"].casefold())) if faculty else None
            if faculty is None:
                errors.append("ไม่พบคณะ")
            if department is None:
                errors.append("ไม่พบสาขาในคณะที่ระบุ")
            resolved["faculty_id"] = faculty["id"] if faculty else None
            resolved["department_id"] = department["id"] if department else None
            if values["user_type"] == "student":
                cohort = cohorts.get((department["id"], values["cohort"].casefold())) if department else None
                group = class_groups.get((cohort["id"], values["class_group"].casefold())) if cohort and values["class_group"] else None
                if cohort is None:
                    errors.append("ไม่พบรุ่นในสาขาที่ระบุ")
                if values["class_group"] and group is None:
                    errors.append("ไม่พบหมู่เรียนในรุ่นที่ระบุ")
                resolved["cohort_id"] = cohort["id"] if cohort else None
                resolved["class_group_id"] = group["id"] if group else None
        if username:
            seen.add(key)
        preview.append(ImportRow(number, values, resolved, tuple(dict.fromkeys(errors))))
    return ImportPreview(tuple(preview))


def import_preview(service: AppService, actor_id: int, preview: ImportPreview) -> tuple[int, list[str]]:
    created = 0
    errors: list[str] = []
    for row in preview.rows:
        if row.errors:
            continue
        try:
            service.save_user(actor_id, **row.resolved)
            created += 1
        except Exception as error:
            errors.append(f"แถว {row.number}: {error}")
    return created, errors
